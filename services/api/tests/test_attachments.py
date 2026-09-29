"""Tests for uploaded spreadsheet and document processing."""

import io
import openpyxl
from app.engine.engine import InsightEngine


def test_uploaded_excel_processing():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Revenue", "Units"])
    ws.append(["North", 500000.0, 100])
    ws.append(["South", 300000.0, 60])
    ws.append(["East", 200000.0, 40])

    buf = io.BytesIO()
    wb.save(buf)

    engine = InsightEngine.fixture()
    attachments = [{"name": "sales_q3.xlsx", "kind": "document", "raw_bytes": buf.getvalue()}]

    # Query total revenue
    env = engine.ask("What is the total revenue in this excel sheet?", attachments=attachments)
    assert env.route == "structured"
    assert "1,000,000" in env.answer or "1000000" in env.answer
    assert env.sql
    assert env.tables
    assert any("sales_q3.xlsx" in c.title for c in env.citations)


def test_uploaded_document_processing():
    doc_text = """
Company Travel Policy 2026:
- Maximum hotel allowance is $250 per night.
- Meal per diem is $75 per day.
- Flight bookings must be made at least 14 days in advance.
    """.strip()

    engine = InsightEngine.fixture()
    attachments = [{"name": "travel_policy.txt", "kind": "document", "raw_bytes": doc_text.encode("utf-8")}]

    env = engine.ask("What is the maximum hotel allowance?", attachments=attachments)
    assert env.route == "unstructured"
    assert "$250" in env.answer
    assert any("travel_policy.txt" in c.title for c in env.citations)
    assert "north" not in env.answer.lower()


def test_uploaded_json_reports_summary():
    from pathlib import Path

    reports_path = Path("d:/capstone/insight-gpt/data/generated/documents/reports.json")
    if not reports_path.exists():
        return
    reports_bytes = reports_path.read_bytes()

    engine = InsightEngine.fixture()
    attachments = [{"name": "reports.json", "kind": "document", "raw_bytes": reports_bytes}]

    env = engine.ask("summarize these files", attachments=attachments)
    assert "*[ *:" not in env.answer
    assert "2025Q3 performance summary" in env.answer
    assert "2026Q2 operations review" in env.answer
    assert len(env.citations) >= 4


def test_multi_file_attachments():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Category", "Sales"])
    ws.append(["Hardware", 5000.0])
    buf = io.BytesIO()
    wb.save(buf)

    doc_text = "Vendor contract notes: SLA response time is 1 hour."

    engine = InsightEngine.fixture()
    attachments = [
        {"name": "sales.xlsx", "kind": "document", "raw_bytes": buf.getvalue()},
        {"name": "contract.txt", "kind": "document", "raw_bytes": doc_text.encode("utf-8")},
    ]

    env = engine.ask("summarize these files", attachments=attachments)
    assert "sales.xlsx" in env.answer
    assert "contract.txt" in env.answer
    assert len(env.citations) >= 2
